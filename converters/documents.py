# converters/documents.py
import os
import shutil
import subprocess
from pathlib import Path
from . import ConversionResult

def _markdown_to_docx_fallback(md_content: str, output_path: Path):
    """Pure-Python Markdown to DOCX parser when Pandoc is unavailable."""
    import docx
    doc = docx.Document()
    
    for line in md_content.splitlines():
        line_s = line.strip()
        if not line_s:
            continue
        if line_s.startswith("### "):
            doc.add_heading(line_s[4:], level=3)
        elif line_s.startswith("## "):
            doc.add_heading(line_s[3:], level=2)
        elif line_s.startswith("# "):
            doc.add_heading(line_s[2:], level=1)
        elif line_s.startswith(("- ", "* ")):
            doc.add_paragraph(line_s[2:], style='List Bullet')
        elif len(line_s) > 3 and line_s[:2].isdigit() and line_s[2:4] in (". ", ") "):
            doc.add_paragraph(line_s[3:], style='List Number')
        elif line_s.startswith("> "):
            p = doc.add_paragraph(line_s[2:])
            p.italic = True
        else:
            doc.add_paragraph(line_s)
            
    doc.save(str(output_path))

def _markdown_to_html_fallback(md_content: str, output_path: Path):
    """Pure-Python basic Markdown to HTML parser when Pandoc is unavailable."""
    html_lines = ["<!DOCTYPE html><html><head><meta charset='utf-8'><style>body{font-family:sans-serif;max-width:800px;margin:2rem auto;padding:0 1rem;line-height:1.6;}</style></head><body>"]
    for line in md_content.splitlines():
        s = line.strip()
        if not s:
            continue
        if s.startswith("### "):
            html_lines.append(f"<h3>{s[4:]}</h3>")
        elif s.startswith("## "):
            html_lines.append(f"<h2>{s[3:]}</h2>")
        elif s.startswith("# "):
            html_lines.append(f"<h1>{s[2:]}</h1>")
        elif s.startswith(("- ", "* ")):
            html_lines.append(f"<li>{s[2:]}</li>")
        else:
            html_lines.append(f"<p>{s}</p>")
    html_lines.append("</body></html>")
    output_path.write_text("\n".join(html_lines), encoding="utf-8")

def convert(input_path: str, output_format: str, output_dir: str) -> ConversionResult:
    input_path = Path(input_path)
    output_dir = Path(output_dir)
    fmt = output_format.lower().strip().lstrip(".")
    
    # Normalize legacy 'doc' to modern 'docx'
    if fmt == 'doc':
        fmt = 'docx'

    output_path = output_dir / f"{input_path.stem}.{fmt}"
    ext = input_path.suffix.lower()

    output_dir.mkdir(parents=True, exist_ok=True)

    try:
        # 1. PDF -> Word (DOCX)
        if ext == '.pdf' and fmt == 'docx':
            try:
                from pdf2docx import Converter
                cv = Converter(str(input_path))
                cv.convert(str(output_path), start=0, end=None)
                cv.close()
                if output_path.exists() and output_path.stat().st_size > 0:
                    return ConversionResult(status='success', output_path=str(output_path), output_size=output_path.stat().st_size)
            except Exception:
                # Fallback engine: PyMuPDF + python-docx
                try:
                    import fitz
                    import docx
                    doc_pdf = fitz.open(str(input_path))
                    doc_word = docx.Document()
                    has_content = False
                    for page_idx in range(len(doc_pdf)):
                        page = doc_pdf[page_idx]
                        text = page.get_text()
                        if text.strip():
                            has_content = True
                            for line in text.splitlines():
                                if line.strip():
                                    doc_word.add_paragraph(line)
                            if page_idx < len(doc_pdf) - 1:
                                doc_word.add_page_break()
                    doc_pdf.close()
                    if not has_content:
                        doc_word.add_paragraph("(No selectable text extracted from PDF)")
                    doc_word.save(str(output_path))
                    if output_path.exists():
                        return ConversionResult(
                            status='success', 
                            output_path=str(output_path), 
                            output_size=output_path.stat().st_size,
                            message="Converted using secondary layout engine"
                        )
                except Exception as fb_err:
                    return ConversionResult(
                        status='error',
                        error=f"PDF to DOCX conversion failed: {fb_err}"
                    )

        # 2. PDF -> Plain Text (TXT)
        if ext == '.pdf' and fmt == 'txt':
            try:
                import fitz
                doc = fitz.open(str(input_path))
                text_pages = [page.get_text() for page in doc]
                doc.close()
                output_path.write_text("\n\n".join(text_pages), encoding="utf-8")
                size = output_path.stat().st_size if output_path.exists() else None
                return ConversionResult(status='success', output_path=str(output_path), output_size=size)
            except ImportError:
                pass

        # 3. DOCX / DOC -> PDF
        if ext in ['.docx', '.doc'] and fmt == 'pdf':
            soffice_cmd = shutil.which("libreoffice") or shutil.which("soffice")
            if soffice_cmd:
                cmd = [soffice_cmd, "--headless", "--convert-to", "pdf", str(input_path), "--outdir", str(output_dir)]
                subprocess.run(cmd, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
                if output_path.exists():
                    return ConversionResult(status='success', output_path=str(output_path), output_size=output_path.stat().st_size)

            try:
                from docx2pdf import convert as docx2pdf_convert
                docx2pdf_convert(str(input_path), str(output_path))
                if output_path.exists():
                    return ConversionResult(status='success', output_path=str(output_path), output_size=output_path.stat().st_size)
            except Exception as docx_err:
                return ConversionResult(
                    status='error', 
                    error=f"DOCX to PDF requires LibreOffice or Microsoft Word. (Run via Docker or install Word/LibreOffice)."
                )

        # 4. DOCX -> TXT
        if ext in ['.docx', '.doc'] and fmt == 'txt':
            import docx
            doc = docx.Document(str(input_path))
            full_text = [p.text for p in doc.paragraphs]
            output_path.write_text("\n".join(full_text), encoding="utf-8")
            return ConversionResult(status='success', output_path=str(output_path), output_size=output_path.stat().st_size)

        # 5. TXT -> DOCX
        if ext == '.txt' and fmt == 'docx':
            import docx
            doc = docx.Document()
            content = input_path.read_text(encoding="utf-8", errors="replace")
            for line in content.splitlines():
                doc.add_paragraph(line)
            doc.save(str(output_path))
            return ConversionResult(status='success', output_path=str(output_path), output_size=output_path.stat().st_size)

        # 6. TXT -> HTML / MD
        if ext == '.txt' and fmt == 'html':
            content = input_path.read_text(encoding="utf-8", errors="replace")
            html_content = f"<!DOCTYPE html><html><head><meta charset='utf-8'><title>{input_path.stem}</title></head><body><pre>{content}</pre></body></html>"
            output_path.write_text(html_content, encoding="utf-8")
            return ConversionResult(status='success', output_path=str(output_path), output_size=output_path.stat().st_size)

        if ext == '.txt' and fmt in ('md', 'markdown'):
            content = input_path.read_text(encoding="utf-8", errors="replace")
            output_path.write_text(content, encoding="utf-8")
            return ConversionResult(status='success', output_path=str(output_path), output_size=output_path.stat().st_size)

        # 7. Markdown (MD) -> DOCX / HTML / PDF
        if ext in ('.md', '.markdown'):
            md_text = input_path.read_text(encoding="utf-8", errors="replace")
            
            # Try Pandoc first
            try:
                import pypandoc
                format_map = {'docx': 'docx', 'html': 'html', 'pdf': 'pdf', 'txt': 'plain'}
                if fmt in format_map:
                    pypandoc.convert_file(str(input_path), to=format_map[fmt], outputfile=str(output_path))
                    if output_path.exists():
                        return ConversionResult(status='success', output_path=str(output_path), output_size=output_path.stat().st_size)
            except Exception:
                # Pandoc missing or failed, trigger pure-python fail-safe
                pass

            # Fail-safe purely in Python:
            if fmt == 'docx':
                _markdown_to_docx_fallback(md_text, output_path)
                if output_path.exists():
                    return ConversionResult(
                        status='success', 
                        output_path=str(output_path), 
                        output_size=output_path.stat().st_size,
                        message="Converted with native Markdown layout engine"
                    )
            elif fmt == 'html':
                _markdown_to_html_fallback(md_text, output_path)
                if output_path.exists():
                    return ConversionResult(status='success', output_path=str(output_path), output_size=output_path.stat().st_size)
            elif fmt == 'txt':
                output_path.write_text(md_text, encoding="utf-8")
                return ConversionResult(status='success', output_path=str(output_path), output_size=output_path.stat().st_size)

        # 8. HTML -> MD / DOCX
        if ext == '.html':
            try:
                import pypandoc
                format_map = {'docx': 'docx', 'md': 'markdown', 'markdown': 'markdown', 'txt': 'plain'}
                if fmt in format_map:
                    pypandoc.convert_file(str(input_path), to=format_map[fmt], outputfile=str(output_path))
                    if output_path.exists():
                        return ConversionResult(status='success', output_path=str(output_path), output_size=output_path.stat().st_size)
            except Exception as pe:
                return ConversionResult(status='error', error=f"HTML conversion requires Pandoc: {pe}")

        # General Pandoc fallback for remaining combos
        format_map = {'docx': 'docx', 'html': 'html', 'md': 'markdown', 'markdown': 'markdown', 'pdf': 'pdf'}
        if fmt in format_map:
            try:
                import pypandoc
                pypandoc.convert_file(str(input_path), to=format_map[fmt], outputfile=str(output_path))
                if output_path.exists():
                    return ConversionResult(status='success', output_path=str(output_path), output_size=output_path.stat().st_size)
            except Exception as pe:
                return ConversionResult(status='error', error=f"Conversion requires Pandoc engine: {pe}")

        return ConversionResult(status='error', error=f"Unsupported document conversion: {ext} -> {fmt}")

    except Exception as e:
        return ConversionResult(status='error', error=str(e))